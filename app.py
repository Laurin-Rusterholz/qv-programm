"""QV-Pruefungsassistent - lokale Backup-App.

Eine einfache, lokale KI-Chat-App fuer die schriftliche Geleitete Fallarbeit
(Kaufleute EFZ, HKB B/C/E). Sie funktioniert wie das bisherige Claude-Projekt
und kann zusaetzlich echte Office-Dateien (Excel, Word, PowerPoint, PDF) erstellen.

Start:  streamlit run app.py      (oder einfach start.bat / start.sh ausfuehren)

Aufbau der Datei:
  1. Konstanten und System-Prompt
  2. Reine Hilfsfunktionen (ohne Streamlit) - dadurch testbar
  3. main() mit der gesamten Streamlit-Oberflaeche
"""

import os
import sys
import base64
import subprocess
from pathlib import Path

import streamlit as st
from dotenv import load_dotenv, set_key

# Das anthropic-SDK wird vorsichtig importiert, damit die App auch dann startet,
# wenn das Paket (noch) fehlt - dann erscheint eine klare Meldung.
try:
    import anthropic
    ANTHROPIC_VERFUEGBAR = True
except Exception:  # pragma: no cover - nur falls das Paket fehlt
    ANTHROPIC_VERFUEGBAR = False


# ==========================================================================
# 1. KONSTANTEN
# ==========================================================================

PROJEKT_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = PROJEKT_DIR / "uploads"
OUTPUT_DIR = PROJEKT_DIR / "outputs"
THEORIE_DIR = PROJEKT_DIR / "theorie"
ENV_PFAD = PROJEKT_DIR / ".env"

# Gut sichtbare Modell-Konstante. Standard: schnell und stark genug.
MODELL = "claude-sonnet-4-6"

# Auswahl in der Seitenleiste: Anzeige-Text -> Modell-ID
MODELL_OPTIONEN = {
    "Schnell (Sonnet 4.6)": "claude-sonnet-4-6",
    "Beste Qualitaet (Opus 4.8)": "claude-opus-4-8",
}

# Preise in US-Dollar pro 1 Million Tokens (Eingabe, Ausgabe) - fuer die grobe
# Kostenanzeige. Stand der offiziellen Anthropic-Preisliste.
PREISE_PRO_MIO = {
    "claude-sonnet-4-6": (3.0, 15.0),
    "claude-opus-4-8": (5.0, 25.0),
}

# Lange Loesungen sollen vollstaendig sein.
MAX_TOKENS = 8000

# API-Version (das SDK sendet diesen Wert; wir setzen ihn zusaetzlich ausdruecklich).
ANTHROPIC_VERSION = "2023-06-01"

# Theorie und angehaengte Office-Dateien werden bei Bedarf gekuerzt.
THEORIE_MAX_ZEICHEN = 24000
DATEI_MAX_ZEICHEN = 30000

# Maximale Korrekturversuche bei fehlerhaftem Datei-Code (Selbstkorrektur).
MAX_FEHLER_VERSUCHE = 3

# Sicherheitsobergrenze fuer die Werkzeug-Schleife (verhindert Endlosschleifen).
MAX_SCHLEIFE = 12

# Welche Endungen werden als Bild nativ an die KI gegeben?
BILD_TYPEN = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}

# MIME-Typen fuer die Download-Knoepfe.
MIME_TYPEN = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".pdf": "application/pdf",
    ".csv": "text/csv",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
}

# Name des Datei-Werkzeugs. Hinweis: Die Anthropic-API erlaubt in Werkzeugnamen
# nur ASCII (a-z, A-Z, 0-9, _, -). Deshalb heisst das Werkzeug technisch
# "python_ausfuehren". Im System-Prompt wird es (wie gewuenscht) als
# "python_ausfuehren" / «python_ausführen» beschrieben - die KI ruft immer das
# tatsaechlich bereitgestellte Werkzeug auf.
TOOL_NAME = "python_ausfuehren"

TOOLS = [
    {
        "name": TOOL_NAME,
        "description": (
            "Fuehrt Python-Code im Projektordner aus, um pruefungstaugliche Dateien "
            "zu erstellen. Dies ist das im System-Prompt genannte Werkzeug "
            "«python_ausführen». Verfuegbar und importierbar sind openpyxl (Excel), "
            "python-docx (import docx, Word), python-pptx (import pptx, PowerPoint), "
            "reportlab (PDF) und pandas. Du kannst zusaetzlich aus den Ordnern "
            "'uploads/' und 'theorie/' lesen (z. B. um eine hochgeladene Vorlage "
            "auszufuellen). Erzeugte Dateien MUESSEN in den Ordner 'outputs/' "
            "geschrieben werden; nutze relative Pfade wie 'outputs/Budget.xlsx'. "
            "Bei Excel immer echte Formeln verwenden (Zellen, die mit '=' beginnen) "
            "und alle Eingabewerte in eigene, klar beschriftete Zellen legen, auf die "
            "die Formeln verweisen. Schreibe nichts ausserhalb des Projektordners."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "code": {
                    "type": "string",
                    "description": (
                        "Der auszufuehrende Python-Code. Speichert die Datei(en) im "
                        "Ordner 'outputs/'."
                    ),
                }
            },
            "required": ["code"],
        },
    }
]


# System-Prompt - woertlich wie vorgegeben.
SYSTEM_PROMPT = """# QV-Prüfungsassistent (KV EFZ, Geleitete Fallarbeit)

## Rolle
Du bist mein QV-Prüfungsassistent für die schriftliche Geleitete Fallarbeit (Kaufleute EFZ, HKB B/C/E). Ich gebe dir zuerst die Ausgangslage, danach einzelne Teilaufgaben. Du lieferst prüfungsfertige, direkt einsetzbare Resultate.

## Ablauf
1. Ich lade die Ausgangslage hoch oder schreibe sie in den Chat (Inhalt: meine Rolle, Firma mit Branche und Tools, Mitarbeitende und Abteilungen, Kernproblem, wichtige Personen, wichtige Daten und Fristen, Zielgruppe, Budget und Einschränkungen). Du antwortest nur mit «Verstanden.».
2. Ich starte eine Teilaufgabe, zum Beispiel mit «Teilaufgabe 1» oder «T1», und gebe dir dazu: das Fach (HKB B, C oder E), die Aufgabenstellung und ob du auf die «Theorie» zugreifen sollst.
3. Du schreibst die Lösung direkt in den Chat.
   - Ich bestimme die Dokumentenart. Ein Dokument erstellst du nur, wenn ich es ausdrücklich verlange.
   - Habe ich nicht alles angegeben oder beantwortet, frag nach.
   - Verstehst du etwas in der Ausgangslage oder Aufgabe nicht, frag kurz nach. Die Ausgangslage kann Schreibfehler enthalten.
4. Du gibst das Resultat aus. Ich stelle Rückfragen oder gebe Anpassungen, und du überarbeitest es.
5. Sobald ich «weiter» schreibe oder eine neue Teilaufgabe gebe, machen wir mit dieser weiter.
6. Ich gebe die Ausgangslage oft nicht komplett auf einmal, sondern nach und nach mehr Kontext. Beziehe immer den gesamten bisherigen Chatverlauf mit ein.

## Regeln
- Prüfungsfertig und direkt einsetzbar, ohne Erklärungen oder Vorreden.
- Antworte zügig und knapp. Wenn ich ein Programm oder Format nenne, setze es sofort um, ohne Alternativen vorzuschlagen oder erneut nachzufragen. Keine Debugging-Hinweise.
- Halte exakte Mengenvorgaben der Aufgabe strikt ein: Wortzahl, Anzahl Sätze, Anzahl Beispiele, Kriterien, Schritte oder Punkte und Zeitangaben. Bei einer Spanne (zum Beispiel «70 bis 90 Wörter») ziele auf die sichere Mitte, nicht an den Rand. Zähle vor der Ausgabe nach und korrigiere, bis die Vorgabe erfüllt ist. Gezählt wird nur der reine Antworttext, ohne Überschriften, Platzhalter oder Bezeichnungen.
- Schweizer Rechtschreibung: «ss» statt «ß».
- Währung CHF, kaufmännisch runden, MwSt 8.1 %.
- E-Mails kurz und prägnant halten: nur das Nötige, keine Füllsätze oder Wiederholungen, wenige kurze Absätze; alle geforderten Inhaltspunkte trotzdem enthalten. E-Mails enden mit «Freundliche Grüsse».
- Bei fremdsprachigen Aufgaben: GER-Niveau C1, formaler Ton.
- Keine Gedankenstriche.
- Immer konkret auf meinen Fall bezogen, nie allgemein.

## Genauigkeit und Form
- Fachliche und rechtliche Angaben exakt, eindeutig und widerspruchsfrei. Was du als korrekt erkennst, muss auch im Endergebnis korrekt stehen.
- Vorgegebene Gliederungen oder Kategorien strikt trennen, jede vollständig und nur mit dem füllen, was wirklich hineingehört (zum Beispiel keine betrieblichen Abläufe bei gesetzlichen Pflichtangaben).
- Stil an die Textsorte anpassen: offizielle Dokumente (Handbuch, Reglement, Vertrag) neutral und sachlich formulieren, nicht in Ich-Form. Innerhalb eines Dokuments einheitlich darstellen, nicht Stichpunkte, Symbole und Zitatform mischen.
- Wenn ein Tool oder Informationsmittel verlangt ist, ein einziges konsistentes Konzept wählen, nicht zwei Mittel vermischen.
- Liegt eine Vorlage oder ein Mockup vor, orientiere dich möglichst exakt daran: sichtbare Gestaltungselemente präzise übernehmen (Farben, Schriften, Linien, Symbole, Seitenformat, Abstände, Reihenfolge) und die vorgegebene Struktur einhalten.
- Berücksichtige Nachhaltigkeit, wo es zum Fall passt: ökologische und soziale Verantwortung, ressourcenschonende und digitale Lösungen, gute ÖV-Anbindung. Im KV-QV ist das oft ein verstecktes Bewertungskriterium.
- Ist nur eine Planung verlangt, erstelle kein fertiges Produkt.

## Dokumente
- Standardmässig schreibst du das Resultat immer direkt in den Chat.
- Ich bestimme die Dokumentenart. Erstelle nur dann ein Dokument, wenn ich es ausdrücklich verlange.
- Verlange ich ein Dokument oder Informationsmittel, ohne das Programm zu nennen, frag zuerst, womit ich es erstellt haben will. Wähle das Programm nie selbst und erstelle nie automatisch HTML.
- Schlage als Programme Microsoft vor (Word, Excel, PowerPoint, Teams, SharePoint, Outlook); für visuelle Informationsmittel wie Poster, Flyer oder Infoblatt eignen sich PowerPoint oder Canva, nie HTML.
- Wenn ich ein Excel verlange, rechnest du darin alles mit echten Formeln, nie mit fest eingetippten Werten.
- Wenn ich ein Word verlange: Tabellen müssen vollständig auf die Seite passen, keine abgeschnittenen Spalten. Bei breiten Tabellen (etwa fünf Spalten oder mehr) Querformat wählen und die Spaltenbreiten so setzen, dass die Tabelle in die nutzbare Seitenbreite passt.

## Projektdateien und Theorie
- Greife nur auf die im Projekt gespeicherten Dateien (Theorie, Vorlagen usw.) zu, wenn ich es ausdrücklich verlange.
- Die gesamte Theorie liegt im Projekt. Nutze sie nur, wenn ich in einer Aufgabe das Wort «Theorie» schreibe. Sonst antwortest du, ohne auf die Projektdateien zuzugreifen.
- Die Theorie begründet die Methode und den Inhalt (zum Beispiel den Diagrammtyp), nie das Dateiformat. Begründe ein Dateiformat nie mit der Theorie.

## Bei HKB E (wenn ich «HKB E» schreibe)
- Informatik und Darstellung zählen mit: technisch und formal korrekt arbeiten, sauber, übersichtlich und professionell formatieren oder visualisieren.
- Budget und Kalkulation: alles mit echten Formeln, und alle Eingabewerte (Preise, Mengen, Sätze, Rabatte, Anzahl Teilnehmende und Referenten, MWST-Satz) in eigenen anpassbaren Zellen, nie fest in die Formel eingebaut. Ändere ich eine Eingabezelle, muss das ganze Budget automatisch stimmen. Aufbau im Detail siehe «Kalkulation und Kostenvoranschlag».
- PowerPoint immer mit einem Folienmaster aufbauen: Layouts, Schriften, Farben und Platzhalter zentral im Folienmaster festlegen, damit die Präsentation einheitlich und professionell ist und sich leicht anpassen lässt.
- In Präsentationen stichwortartig schreiben, keine Fliesstexte: kurze, prägnante Stichpunkte statt ausformulierter Sätze auf den Folien.
- Allgemeine HKB-E-Aufgaben: so gestalten, dass das Resultat wiederverwendbar ist, zum Beispiel als Vorlage für Folgejahre.
- Auf Details achten: fehlende oder ungenaue Details geben im QV Punktabzug.
- Das Ergebnis muss nachvollziehbar sein. Markiere farbig, was verlangt wird.

## Worauf bei den typischen Aufgabentypen zu achten ist
- Ablage- und Dokumentationsstruktur: passende digitale Arbeitsumgebung begründen; klare Struktur mit Haupt- und Unterkategorien, für Folgejahre wiederverwendbar.
- Datenschutz und Datensicherheit: konkrete Beispiele schützenswerter Daten oder Abläufe mit Fallbezug, je mit Begründung und konkreter Massnahme.
- Projektplanung: jede Aufgabe einzeln auflisten (auch einzelne Event-Aktivitäten wie Konzerte und Shows sowie Auf- und Abbau der Technik), nichts zusammenfassen; logische Reihenfolge beachten (Technik und Infrastruktur vor Veranstaltungsbeginn aufbauen); Beginn und Ende der Raumnutzung als eigene Ereignisse; je mit Start, Ende, Verantwortlich und Status; Meilensteine kennzeichnen; alles innerhalb des vorgegebenen Zeitfensters; übersichtlich als Gantt-Diagramm.
- Texte überarbeiten oder gliedern (zum Beispiel Reglement, Handbuch, Lehrvertrag): nur die verlangten Kategorien, jede sauber abgegrenzt und vollständig; fachlich und rechtlich korrekt; neutrale, einheitliche Formulierung.
- Ablaufprozess oder Pflichtenheft: jeden Schritt oder jede Aufgabe einzeln, mit Verantwortlichkeit (Zentrale oder Standort) und wo nötig Entscheidungsinstanz; lückenlos und in logischer Reihenfolge, inklusive Abschluss (zum Beispiel Check-out, Verabschiedung, Feedback vor Ort, wo relevant Zahlungs- oder Stornoprozess); nah an einer gegebenen Vorlage.
- Konfliktlösung: Vorgehen Schritt für Schritt mit Begründung, gestützt auf ein anerkanntes Modell (zum Beispiel gewaltfreie Kommunikation: Wahrnehmung, Gefühl, Bedürfnis, Bitte); dazu Massnahmen, damit das Problem künftig nicht mehr auftritt.
- Recherche und Entscheidung: sinnvolle Kriterien ableiten, mehrere Optionen prüfen, Vergleich in einer übersichtlichen Tabelle, Entscheid plausibel und empfängergerecht begründen.
- Kommunikationsprozess und To-do-Liste: Was, Wer, Bis wann; alle Zielgruppen berücksichtigen; realistische Verantwortlichkeiten und Fristen.
- Texte wie News, Posts oder E-Mails: zielgruppengerecht und sprachlich korrekt, da Ausdruck, Rechtschreibung und Grammatik in die Bewertung einfliessen.
- Werbe- oder Informationsmaterial (Infoblatt, Flyer, Poster, Newsbeitrag): klare visuelle Gliederung in Abschnitte (Titel, Einstieg, Inhalt oder Programm, Anmeldung); immer eine klare Handlungsaufforderung mit Anmelde- oder Buchungsinfo (Link und, wenn möglich, QR-Code); adressatengerecht und vollständig.
- Kalkulation und Kostenvoranschlag: fixe und variable Kosten in getrennten Abschnitten; zusammengesetzte Posten in ihre Bestandteile aufteilen (zum Beispiel Honorar in Basis, Zuschlag und Standard); alle Eingabewerte (Preise, Mengen, Sätze, Rabatte, Anzahl Teilnehmende und Referenten) in eigenen anpassbaren Zellen, nie fest in die Formel; Spalte Preis pro Person oder Stück; Total exkl. MWST, MWST 8.1 % und Total inkl. MWST; kaufmännisch gerundet.
- Verträge anpassen: alle Angaben übertragen, fehlende Felder farbig markieren, geforderte Klauseln wie Zahlungsfristen korrekt einfügen.
- Fremdsprachige Korrespondenz: Niveau C1 und formal, alle geforderten Inhaltspunkte enthalten, übliche Geschäftsformulierungen.

## Technischer Zusatz für diese App
- Dokumente werden erstellt, indem du das Werkzeug `python_ausführen` aufrufst und darin Python mit openpyxl, python-docx, python-pptx oder reportlab schreibst, das die Datei in den Ordner `outputs/` speichert.
- Hochgeladene Dateien liegen im Ordner `uploads/`. Aus diesem Ordner kannst du auch eine vorgegebene Vorlage lesen und ausfüllen.
- Theorie-Dateien liegen im Ordner `theorie/` und werden dir nur dann mitgegeben, wenn meine Nachricht das Wort «Theorie» enthält.
- Wenn du eine Datei erstellt hast, sag mir nur kurz, dass sie bereit ist. Die Oberfläche zeigt automatisch einen Download-Knopf.
"""


# ==========================================================================
# 2. REINE HILFSFUNKTIONEN (ohne Streamlit, daher testbar)
# ==========================================================================

def ordner_anlegen():
    """Legt uploads/, outputs/ und theorie/ an, falls sie fehlen."""
    for ordner in (UPLOAD_DIR, OUTPUT_DIR, THEORIE_DIR):
        ordner.mkdir(parents=True, exist_ok=True)


def sicherer_dateiname(name):
    """Gibt nur den reinen Dateinamen zurueck (kein Pfad, keine Trickserei)."""
    reiner_name = os.path.basename(str(name)).strip()
    reiner_name = reiner_name.replace("\\", "_").replace("/", "_")
    return reiner_name or "datei"


def mime_fuer(dateiname):
    """Liefert den passenden MIME-Typ fuer einen Download-Knopf."""
    return MIME_TYPEN.get(Path(dateiname).suffix.lower(), "application/octet-stream")


def extrahiere_pdf_text(pfad):
    """Liest den Text aus einer PDF-Datei (fuer den Theorie-Ordner)."""
    from pypdf import PdfReader

    reader = PdfReader(str(pfad))
    seiten = []
    for seite in reader.pages:
        seiten.append(seite.extract_text() or "")
    return "\n".join(seiten)


def extrahiere_docx_text(pfad):
    """Liest Absaetze und Tabellen aus einer Word-Datei."""
    import docx

    dokument = docx.Document(str(pfad))
    teile = [absatz.text for absatz in dokument.paragraphs]
    for tabelle in dokument.tables:
        for zeile in tabelle.rows:
            zellen = [zelle.text for zelle in zeile.cells]
            teile.append(" | ".join(zellen))
    return "\n".join(teile)


def extrahiere_xlsx_text(pfad):
    """Liest die Zellwerte aus einer Excel-Datei (berechnete Werte)."""
    import openpyxl

    workbook = openpyxl.load_workbook(str(pfad), data_only=True, read_only=True)
    teile = []
    for blatt in workbook.worksheets:
        teile.append(f"# Blatt: {blatt.title}")
        for zeile in blatt.iter_rows(values_only=True):
            werte = [str(zelle) for zelle in zeile if zelle is not None]
            if werte:
                teile.append(" | ".join(werte))
    workbook.close()
    return "\n".join(teile)


def extrahiere_text_datei(pfad):
    """Liest eine reine Textdatei (txt, csv, md, json)."""
    return Path(pfad).read_text(encoding="utf-8", errors="replace")


def extrahiere_datei_text(pfad):
    """Holt den Text aus einer Datei je nach Endung. Bei Fehlern: None."""
    pfad = Path(pfad)
    suffix = pfad.suffix.lower()
    try:
        if suffix == ".pdf":
            return extrahiere_pdf_text(pfad)
        if suffix == ".docx":
            return extrahiere_docx_text(pfad)
        if suffix in (".xlsx", ".xlsm"):
            return extrahiere_xlsx_text(pfad)
        if suffix in (".txt", ".csv", ".md", ".json", ".text"):
            return extrahiere_text_datei(pfad)
        # Unbekannte Endung: vorsichtig als Text versuchen.
        return extrahiere_text_datei(pfad)
    except Exception:
        return None


def lade_theorie(max_zeichen=THEORIE_MAX_ZEICHEN):
    """Liest den Text aller Dateien im Ordner theorie/ und kuerzt bei Bedarf."""
    if not THEORIE_DIR.exists():
        return ""
    teile = []
    for pfad in sorted(THEORIE_DIR.iterdir()):
        if not pfad.is_file() or pfad.name.startswith("."):
            continue
        text = extrahiere_datei_text(pfad)
        if not text or not text.strip():
            continue
        teile.append(f"===== {pfad.name} =====\n{text.strip()}")
        if sum(len(t) for t in teile) > max_zeichen:
            break
    ergebnis = "\n\n".join(teile)
    if len(ergebnis) > max_zeichen:
        ergebnis = ergebnis[:max_zeichen] + "\n[... Theorie hier gekuerzt ...]"
    return ergebnis


def datei_zu_content_bloecke(pfad):
    """Wandelt eine hochgeladene Datei in Inhalts-Bloecke fuer die KI um.

    PDF und Bilder werden nativ uebergeben, Office-/Textdateien als Text.
    """
    pfad = Path(pfad)
    name = pfad.name
    suffix = pfad.suffix.lower()

    if suffix == ".pdf":
        daten = base64.standard_b64encode(pfad.read_bytes()).decode("ascii")
        return [
            {"type": "text", "text": f"[Hochgeladene Datei: {name}]"},
            {
                "type": "document",
                "source": {
                    "type": "base64",
                    "media_type": "application/pdf",
                    "data": daten,
                },
            },
        ]

    if suffix in BILD_TYPEN:
        daten = base64.standard_b64encode(pfad.read_bytes()).decode("ascii")
        return [
            {"type": "text", "text": f"[Hochgeladenes Bild: {name}]"},
            {
                "type": "image",
                "source": {
                    "type": "base64",
                    "media_type": BILD_TYPEN[suffix],
                    "data": daten,
                },
            },
        ]

    # Word, Excel, TXT usw. -> Text extrahieren.
    text = extrahiere_datei_text(pfad)
    if text is None:
        return [
            {
                "type": "text",
                "text": f"[Hochgeladene Datei {name}: konnte nicht gelesen werden.]",
            }
        ]
    if len(text) > DATEI_MAX_ZEICHEN:
        text = text[:DATEI_MAX_ZEICHEN] + "\n[... gekuerzt ...]"
    return [{"type": "text", "text": f"[Inhalt der hochgeladenen Datei {name}]:\n{text}"}]


def baue_user_inhalt(text, datei_pfade, theorie_text):
    """Baut den Inhalt einer Nutzer-Nachricht (Theorie, Anhaenge, Text)."""
    inhalt = []
    if theorie_text:
        inhalt.append(
            {
                "type": "text",
                "text": (
                    "[THEORIE-KONTEXT aus dem Ordner theorie/ - nur mitgegeben, weil "
                    "die Nachricht das Wort «Theorie» enthaelt. Nutze ihn fuer Methode "
                    "und Inhalt, nie zur Begruendung eines Dateiformats.]\n\n"
                    + theorie_text
                ),
            }
        )
    for pfad in datei_pfade:
        try:
            inhalt.extend(datei_zu_content_bloecke(pfad))
        except Exception as fehler:
            inhalt.append(
                {
                    "type": "text",
                    "text": f"[Datei {Path(pfad).name} konnte nicht verarbeitet werden: {fehler}]",
                }
            )
    if text and text.strip():
        inhalt.append({"type": "text", "text": text})
    if not inhalt:
        inhalt.append({"type": "text", "text": "(leere Nachricht)"})
    return inhalt


def liste_outputs():
    """Momentaufnahme des Ordners outputs/: {Dateiname: Aenderungszeit}."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    eintraege = {}
    for pfad in OUTPUT_DIR.iterdir():
        if pfad.is_file() and not pfad.name.startswith("."):
            try:
                eintraege[pfad.name] = pfad.stat().st_mtime_ns
            except OSError:
                pass
    return eintraege


def python_ausfuehren(code, timeout=120):
    """Fuehrt KI-Code im Projektordner aus und meldet neue Dateien zurueck.

    Rueckgabe (dict):
      fehler:        None bei Erfolg, sonst die vollstaendige Fehlermeldung (Text)
      ausgabe:       die Standardausgabe (stdout) des Codes
      neue_dateien:  Liste neuer/geaenderter Dateinamen im Ordner outputs/
    """
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    vorher = liste_outputs()

    try:
        # Der Code wird ueber stdin an einen frischen Python-Prozess gegeben.
        # cwd = Projektordner, damit relative Pfade (outputs/, uploads/, theorie/)
        # funktionieren. So kann ein Fehler die App nicht zum Absturz bringen.
        ergebnis = subprocess.run(
            [sys.executable, "-"],
            input=code,
            cwd=str(PROJEKT_DIR),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
        )
    except subprocess.TimeoutExpired:
        return {
            "fehler": f"Zeitueberschreitung: Der Code lief laenger als {timeout} Sekunden.",
            "ausgabe": "",
            "neue_dateien": [],
        }
    except Exception as fehler:  # pragma: no cover - sehr seltener Sonderfall
        return {
            "fehler": f"Konnte den Code nicht starten: {fehler}",
            "ausgabe": "",
            "neue_dateien": [],
        }

    nachher = liste_outputs()
    neue_dateien = sorted(
        name
        for name, zeit in nachher.items()
        if name not in vorher or vorher.get(name) != zeit
    )

    if ergebnis.returncode != 0:
        fehlertext = (ergebnis.stderr or "").strip() or "Unbekannter Fehler (kein stderr)."
        return {
            "fehler": fehlertext,
            "ausgabe": ergebnis.stdout or "",
            "neue_dateien": neue_dateien,
        }

    return {
        "fehler": None,
        "ausgabe": ergebnis.stdout or "",
        "neue_dateien": neue_dateien,
    }


def block_zu_dict(block):
    """Wandelt einen Antwort-Block des SDK in ein einfaches dict (zum Zurueckschicken)."""
    if hasattr(block, "model_dump"):
        return block.model_dump(exclude_none=True, mode="json")
    return block


def text_aus_bloecken(bloecke):
    """Sammelt den reinen Text aus einer Liste von Antwort-Bloecken."""
    teile = []
    for block in bloecke:
        if getattr(block, "type", None) == "text":
            teile.append(block.text)
    return "\n".join(teile).strip()


def geschaetzte_kosten(modell, input_tokens, output_tokens):
    """Grobe Kostenschaetzung in US-Dollar fuer einen Verbrauch."""
    ein, aus = PREISE_PRO_MIO.get(modell, (3.0, 15.0))
    return (input_tokens / 1_000_000) * ein + (output_tokens / 1_000_000) * aus


def fuehre_konversation(client, modell, messages, status_callback=None):
    """Fuehrt die Werkzeug-Schleife (inkl. Selbstkorrektur) aus.

    messages wird dabei direkt erweitert (voller Verlauf bleibt erhalten).
    Rueckgabe: (antwort_text, neue_dateien, verbrauch) mit
    verbrauch = {"input": int, "output": int}.
    """
    def melde(text):
        if status_callback:
            try:
                status_callback(text)
            except Exception:
                pass

    def zaehle(antwort):
        nutzung = getattr(antwort, "usage", None)
        if not nutzung:
            return
        verbrauch["input"] += (
            (getattr(nutzung, "input_tokens", 0) or 0)
            + (getattr(nutzung, "cache_read_input_tokens", 0) or 0)
            + (getattr(nutzung, "cache_creation_input_tokens", 0) or 0)
        )
        verbrauch["output"] += getattr(nutzung, "output_tokens", 0) or 0

    neue_dateien = []
    fehler_anzahl = 0
    letzter_text = ""
    verbrauch = {"input": 0, "output": 0}

    system_param = [
        {
            "type": "text",
            "text": SYSTEM_PROMPT,
            "cache_control": {"type": "ephemeral"},
        }
    ]

    for _ in range(MAX_SCHLEIFE):
        melde("Die KI denkt nach ...")
        antwort = client.messages.create(
            model=modell,
            max_tokens=MAX_TOKENS,
            system=system_param,
            tools=TOOLS,
            messages=messages,
        )
        zaehle(antwort)

        # Antwort der KI in den Verlauf aufnehmen.
        messages.append(
            {"role": "assistant", "content": [block_zu_dict(b) for b in antwort.content]}
        )
        text = text_aus_bloecken(antwort.content)
        if text:
            letzter_text = text

        if antwort.stop_reason != "tool_use":
            return letzter_text, neue_dateien, verbrauch

        # Die KI moechte das Datei-Werkzeug benutzen.
        tool_ergebnisse = []
        for block in antwort.content:
            if getattr(block, "type", None) != "tool_use":
                continue

            if block.name != TOOL_NAME:
                tool_ergebnisse.append(
                    {
                        "type": "tool_result",
                        "tool_use_id": block.id,
                        "content": "Unbekanntes Werkzeug.",
                        "is_error": True,
                    }
                )
                continue

            eingabe = block.input if isinstance(block.input, dict) else {}
            code = eingabe.get("code", "")
            melde("Die KI erstellt eine Datei ...")
            ergebnis = python_ausfuehren(code)

            if ergebnis["fehler"]:
                fehler_anzahl += 1
                inhalt = "FEHLER bei der Ausfuehrung:\n" + ergebnis["fehler"]
                if ergebnis["ausgabe"].strip():
                    inhalt += "\n\nBisherige Ausgabe:\n" + ergebnis["ausgabe"].strip()
                if fehler_anzahl >= MAX_FEHLER_VERSUCHE:
                    inhalt += (
                        "\n\nDies war der letzte Korrekturversuch. Schreibe keinen "
                        "weiteren Code, sondern erklaere mir kurz in einem Satz, "
                        "woran es liegt."
                    )
                else:
                    inhalt += (
                        "\n\nBitte korrigiere den Code und rufe das Werkzeug erneut auf."
                    )
                tool_ergebnisse.append(
                    {
                        "type": "tool_result",
                        "tool_use_id": block.id,
                        "content": inhalt,
                        "is_error": True,
                    }
                )
            else:
                for name in ergebnis["neue_dateien"]:
                    if name not in neue_dateien:
                        neue_dateien.append(name)
                ausgabe = ergebnis["ausgabe"].strip() or "(keine Textausgabe)"
                dateien = ", ".join(ergebnis["neue_dateien"]) or "(keine neuen Dateien)"
                inhalt = (
                    "Ausfuehrung erfolgreich.\n"
                    f"Standardausgabe:\n{ausgabe}\n"
                    f"Neu erstellte Dateien in outputs/: {dateien}"
                )
                tool_ergebnisse.append(
                    {
                        "type": "tool_result",
                        "tool_use_id": block.id,
                        "content": inhalt,
                        "is_error": False,
                    }
                )

        messages.append({"role": "user", "content": tool_ergebnisse})

        # Sicherheitsausstieg: zu viele Fehler -> noch eine Antwort einholen, dann Schluss.
        if fehler_anzahl >= MAX_FEHLER_VERSUCHE:
            melde("Die KI fasst zusammen ...")
            try:
                abschluss = client.messages.create(
                    model=modell,
                    max_tokens=MAX_TOKENS,
                    system=system_param,
                    tools=TOOLS,
                    messages=messages,
                )
                zaehle(abschluss)
                messages.append(
                    {
                        "role": "assistant",
                        "content": [block_zu_dict(b) for b in abschluss.content],
                    }
                )
                text = text_aus_bloecken(abschluss.content)
                if text:
                    letzter_text = text
            except Exception:
                pass
            if not letzter_text:
                letzter_text = (
                    "Die Datei konnte nach mehreren Versuchen nicht erstellt werden. "
                    "Bitte formuliere die Aufgabe etwas anders oder versuche es erneut."
                )
            return letzter_text, neue_dateien, verbrauch

    # Sicherheitsobergrenze erreicht.
    if not letzter_text:
        letzter_text = "Die KI hat keine abschliessende Antwort geliefert. Bitte erneut versuchen."
    return letzter_text, neue_dateien, verbrauch


def bereinige_verlauf(messages):
    """Entfernt eine letzte Assistenten-Nachricht mit offenem Werkzeugaufruf.

    Schutz nach einem Fehler: Ein tool_use ohne zugehoeriges tool_result wuerde
    die naechste Anfrage ungueltig machen.
    """
    while messages and messages[-1].get("role") == "assistant":
        inhalt = messages[-1].get("content")
        hat_tool_use = isinstance(inhalt, list) and any(
            isinstance(b, dict) and b.get("type") == "tool_use" for b in inhalt
        )
        if hat_tool_use:
            messages.pop()
        else:
            break


def teste_verbindung(api_key, modell):
    """Sendet eine winzige Testanfrage. Rueckgabe: (erfolg, meldung)."""
    if not ANTHROPIC_VERFUEGBAR:
        return False, "Das Paket 'anthropic' ist nicht installiert. Bitte App neu starten."
    try:
        client = anthropic.Anthropic(
            api_key=api_key,
            default_headers={"anthropic-version": ANTHROPIC_VERSION},
        )
        client.messages.create(
            model=modell,
            max_tokens=16,
            messages=[{"role": "user", "content": "Antworte nur mit: OK"}],
        )
        return True, "Verbindung erfolgreich. Die KI hat geantwortet."
    except anthropic.AuthenticationError:
        return False, "Der API-Schluessel ist ungueltig. Bitte pruefe ihn."
    except anthropic.APIConnectionError:
        return False, "Keine Verbindung zur KI. Bitte pruefe deine Internetverbindung."
    except anthropic.RateLimitError:
        return False, "Zu viele Anfragen (Rate-Limit). Bitte kurz warten und erneut testen."
    except anthropic.APIStatusError as fehler:
        return (
            False,
            f"Die KI meldet einen Fehler (Code {fehler.status_code}). "
            "Moeglicherweise fehlt API-Guthaben auf deinem Konto.",
        )
    except Exception as fehler:
        return False, f"Unerwarteter Fehler: {fehler}"


# ==========================================================================
# 3. STREAMLIT-OBERFLAECHE
# ==========================================================================

def zeige_download_knopf(dateiname, key):
    """Zeigt einen Download-Knopf fuer eine Datei aus dem Ordner outputs/."""
    pfad = OUTPUT_DIR / dateiname
    if not pfad.exists():
        st.caption(f"Datei {dateiname} ist nicht mehr vorhanden.")
        return
    try:
        daten = pfad.read_bytes()
    except OSError:
        st.caption(f"Datei {dateiname} konnte nicht gelesen werden.")
        return
    st.download_button(
        label=f"⬇️  {dateiname}  herunterladen",
        data=daten,
        file_name=dateiname,
        mime=mime_fuer(dateiname),
        key=key,
    )


def _secret_oder_none(name):
    """Liest einen Wert aus st.secrets, ohne zu crashen, wenn keine secrets.toml existiert."""
    try:
        if name in st.secrets:
            return st.secrets[name]
    except Exception:
        return None
    return None


def hole_api_key_aus_quellen():
    """API-Schluessel aus st.secrets (Cloud) oder Umgebung/.env (lokal)."""
    aus_secret = _secret_oder_none("ANTHROPIC_API_KEY")
    if aus_secret:
        return str(aus_secret)
    return os.environ.get("ANTHROPIC_API_KEY", "")


def pruefe_passwort():
    """Optionaler Passwortschutz fuer den Online-Betrieb.

    Aktiv, sobald APP_PASSWORT in den Secrets oder der Umgebung gesetzt ist.
    Ohne gesetztes Passwort ist die App offen (lokaler Betrieb).
    Rueckgabe True, wenn der Zugang frei ist.
    """
    erwartet = _secret_oder_none("APP_PASSWORT") or os.environ.get("APP_PASSWORT")
    if not erwartet:
        return True
    if st.session_state.get("_zugang_ok"):
        return True

    st.title("🔒 QV-Pruefungsassistent")
    st.write("Diese App ist passwortgeschuetzt. Bitte Passwort eingeben.")
    pw = st.text_input("Passwort", type="password")
    if st.button("Anmelden"):
        if pw == str(erwartet):
            st.session_state["_zugang_ok"] = True
            st.rerun()
        else:
            st.error("Falsches Passwort.")
    return False


def _dateien_im_ordner(ordner, ausnahmen=()):
    """Sortierte Liste sichtbarer Dateien in einem Ordner."""
    if not ordner.exists():
        return []
    return [
        p
        for p in sorted(ordner.iterdir())
        if p.is_file() and not p.name.startswith(".") and p.name not in ausnahmen
    ]


def verwalte_theorie():
    """Sidebar-Bereich: Theorie-Dateien hochladen, ansehen, loeschen."""
    with st.sidebar.expander("📚 Theorie-Dateien", expanded=False):
        st.caption(
            "Wird nur genutzt, wenn deine Nachricht das Wort «Theorie» enthaelt."
        )
        neu = st.file_uploader(
            "Theorie hinzufuegen",
            accept_multiple_files=True,
            key="theorie_upload",
            label_visibility="collapsed",
        )
        if neu:
            for datei in neu:
                ziel = THEORIE_DIR / sicherer_dateiname(datei.name)
                daten = datei.getvalue()
                try:
                    if (not ziel.exists()) or ziel.stat().st_size != len(daten):
                        ziel.write_bytes(daten)
                except OSError:
                    pass
        dateien = _dateien_im_ordner(THEORIE_DIR, ausnahmen=("LIESMICH.txt",))
        if dateien:
            for i, pfad in enumerate(dateien):
                c1, c2 = st.columns([5, 1])
                c1.write(pfad.name)
                if c2.button("🗑", key=f"theo_del_{i}", help="loeschen"):
                    pfad.unlink(missing_ok=True)
                    st.rerun()
        else:
            st.caption("Noch keine Theorie-Dateien.")


def verwalte_outputs():
    """Sidebar-Bereich: alle erstellten Dateien auflisten, herunterladen, loeschen."""
    with st.sidebar.expander("📂 Erstellte Dateien", expanded=False):
        dateien = _dateien_im_ordner(OUTPUT_DIR)
        if not dateien:
            st.caption("Noch keine Dateien erstellt.")
            return
        for i, pfad in enumerate(dateien):
            zeige_download_knopf(pfad.name, key=f"side_dl_{i}")
            if st.button("🗑 loeschen", key=f"out_del_{i}"):
                pfad.unlink(missing_ok=True)
                st.rerun()


def verwalte_uploads():
    """Sidebar-Bereich: hochgeladene Dateien auflisten und loeschen."""
    with st.sidebar.expander("📁 Hochgeladene Dateien", expanded=False):
        dateien = _dateien_im_ordner(UPLOAD_DIR)
        if not dateien:
            st.caption("Noch nichts hochgeladen.")
            return
        for i, pfad in enumerate(dateien):
            c1, c2 = st.columns([5, 1])
            c1.write(pfad.name)
            if c2.button("🗑", key=f"up_del_{i}", help="loeschen"):
                pfad.unlink(missing_ok=True)
                st.session_state.angehaengt.discard(pfad.name)
                st.rerun()


def starte_session_state():
    """Legt die noetigen Werte in st.session_state an."""
    if "messages" not in st.session_state:
        st.session_state.messages = []          # voller Verlauf fuer die API
    if "anzeige" not in st.session_state:
        st.session_state.anzeige = []           # Verlauf fuer die Anzeige
    if "angehaengt" not in st.session_state:
        st.session_state.angehaengt = set()     # bereits gesendete Datei-Namen
    if "api_key" not in st.session_state:
        load_dotenv(ENV_PFAD)
        st.session_state.api_key = hole_api_key_aus_quellen()
    if "modell" not in st.session_state:
        st.session_state.modell = MODELL
    if "kosten_total" not in st.session_state:
        st.session_state.kosten_total = 0.0
    if "tokens_total" not in st.session_state:
        st.session_state.tokens_total = 0


def baue_seitenleiste():
    """Baut die Seitenleiste und gibt nichts zurueck (nutzt session_state)."""
    with st.sidebar:
        st.header("Einstellungen")

        # Modellauswahl
        labels = list(MODELL_OPTIONEN.keys())
        # Aktuelles Modell als Vorauswahl bestimmen.
        aktueller_index = 0
        for i, label in enumerate(labels):
            if MODELL_OPTIONEN[label] == st.session_state.modell:
                aktueller_index = i
                break
        gewaehlt = st.radio("Modell", labels, index=aktueller_index)
        st.session_state.modell = MODELL_OPTIONEN[gewaehlt]

        st.divider()

        # API-Schluessel
        eingabe = st.text_input(
            "API-Schluessel",
            value=st.session_state.api_key,
            type="password",
            help=(
                "Bekommst du auf console.anthropic.com (API-Guthaben noetig). "
                "Online kannst du ihn auch dauerhaft in den App-Secrets hinterlegen "
                "(ANTHROPIC_API_KEY)."
            ),
        )
        if eingabe != st.session_state.api_key:
            st.session_state.api_key = eingabe

        spalte1, spalte2 = st.columns(2)
        with spalte1:
            if st.button("Verbindung testen", use_container_width=True):
                if not st.session_state.api_key:
                    st.error("Bitte zuerst einen Schluessel eingeben.")
                else:
                    with st.spinner("Teste Verbindung ..."):
                        erfolg, meldung = teste_verbindung(
                            st.session_state.api_key, st.session_state.modell
                        )
                    if erfolg:
                        st.success(meldung)
                    else:
                        st.error(meldung)
        with spalte2:
            if st.button("In .env speichern", use_container_width=True):
                if not st.session_state.api_key:
                    st.error("Kein Schluessel eingegeben.")
                else:
                    try:
                        set_key(str(ENV_PFAD), "ANTHROPIC_API_KEY", st.session_state.api_key)
                        st.success("In .env gespeichert.")
                    except Exception as fehler:
                        st.error(f"Speichern fehlgeschlagen: {fehler}")

        st.divider()

        if st.button("Neuer Chat (Verlauf leeren)", use_container_width=True):
            st.session_state.messages = []
            st.session_state.anzeige = []
            st.session_state.angehaengt = set()
            st.rerun()

        st.divider()
        st.markdown(
            "**Verbrauch diese Sitzung**  \n"
            f"ca. **US$ {st.session_state.get('kosten_total', 0.0):.3f}** · "
            f"{st.session_state.get('tokens_total', 0):,} Tokens"
        )

        st.divider()
        st.markdown(
            "**So bedienst du die App**\n\n"
            "1. Ausgangslage eintippen oder als Datei anhaengen. Die KI antwortet nur "
            "mit «Verstanden.».\n"
            "2. Teilaufgabe starten, z. B. «T1», mit Fach (HKB B/C/E) und Aufgabe.\n"
            "3. Ein Dokument entsteht nur, wenn du es ausdruecklich verlangst "
            "(z. B. «als Excel»).\n"
            "4. Rueckfragen stellen oder mit «weiter» fortfahren.\n\n"
            "**Dateien anhaengen:** mit dem Bueroklammer-Symbol 📎 unten im Eingabefeld "
            "(auch **mehrere auf einmal**).\n\n"
            "Das Wort **«Theorie»** in deiner Nachricht zieht die Dateien aus dem "
            "Ordner `theorie/` hinzu."
        )

    # Eigene Sidebar-Bereiche zum Verwalten der Dateien.
    verwalte_theorie()
    verwalte_outputs()
    verwalte_uploads()


def entpacke_chat_eingabe(eingabe):
    """Zerlegt die Chat-Eingabe in (Text, Liste hochgeladener Dateien).

    st.chat_input liefert mit accept_file ein Objekt mit .text und .files,
    ohne accept_file einen reinen String. Beides wird hier abgefangen.
    """
    if isinstance(eingabe, str):
        return eingabe, []
    text = getattr(eingabe, "text", "") or ""
    dateien = list(getattr(eingabe, "files", None) or [])
    return text, dateien


def speichere_uploads(dateien):
    """Speichert hochgeladene Dateien in uploads/ und gibt ihre Pfade zurueck."""
    pfade = []
    for datei in dateien:
        ziel = UPLOAD_DIR / sicherer_dateiname(datei.name)
        try:
            ziel.write_bytes(datei.getvalue())
        except OSError:
            continue
        pfade.append(ziel)
    return pfade


def zeige_verlauf():
    """Rendert den bisherigen Chatverlauf samt Download-Knoepfen."""
    for i, nachricht in enumerate(st.session_state.anzeige):
        with st.chat_message(nachricht["role"]):
            if nachricht.get("text"):
                st.markdown(nachricht["text"])
            for j, dateiname in enumerate(nachricht.get("dateien", [])):
                zeige_download_knopf(dateiname, key=f"dl_{i}_{j}")
            if nachricht.get("info"):
                st.caption(nachricht["info"])


def fehlermeldung_fuer(fehler):
    """Macht aus einer Ausnahme eine verstaendliche deutsche Meldung."""
    if ANTHROPIC_VERFUEGBAR and isinstance(fehler, anthropic.AuthenticationError):
        return "Der API-Schluessel ist ungueltig. Bitte links korrigieren."
    if ANTHROPIC_VERFUEGBAR and isinstance(fehler, anthropic.APIConnectionError):
        return "Keine Internetverbindung zur KI. Bitte Internet pruefen und erneut senden."
    if ANTHROPIC_VERFUEGBAR and isinstance(fehler, anthropic.RateLimitError):
        return "Zu viele Anfragen (Rate-Limit). Bitte kurz warten und erneut senden."
    if ANTHROPIC_VERFUEGBAR and isinstance(fehler, anthropic.APIStatusError):
        return (
            f"Die KI meldet einen Fehler (Code {fehler.status_code}). "
            "Moeglicherweise fehlt API-Guthaben."
        )
    return f"Unerwarteter Fehler: {fehler}"


def behandle_eingabe(frage, datei_pfade):
    """Verarbeitet eine neue Nutzer-Nachricht: anhaengen, senden, anzeigen."""
    # Theorie nur laden, wenn das Wort vorkommt.
    theorie_text = lade_theorie() if "theorie" in frage.lower() else ""

    user_inhalt = baue_user_inhalt(frage, datei_pfade, theorie_text)
    st.session_state.messages.append({"role": "user", "content": user_inhalt})

    # Anzeige-Text der Nutzer-Nachricht.
    anzeige_text = frage if frage.strip() else "_(nur Datei(en) gesendet)_"
    if datei_pfade:
        anzeige_text += (
            "\n\n*📎 angehaengt: " + ", ".join(p.name for p in datei_pfade) + "*"
        )
    if theorie_text:
        anzeige_text += "\n\n*Theorie-Kontext wurde mitgegeben*"
    st.session_state.anzeige.append({"role": "user", "text": anzeige_text, "dateien": []})

    with st.chat_message("user"):
        st.markdown(anzeige_text)

    # Antwort der KI holen.
    with st.chat_message("assistant"):
        status = st.status("Die KI denkt nach ...", expanded=False)
        platzhalter = st.empty()
        verbrauch = {"input": 0, "output": 0}
        try:
            client = anthropic.Anthropic(
                api_key=st.session_state.api_key,
                default_headers={"anthropic-version": ANTHROPIC_VERSION},
            )
            antwort_text, neue_dateien, verbrauch = fuehre_konversation(
                client,
                st.session_state.modell,
                st.session_state.messages,
                status_callback=lambda t: status.update(label=t),
            )
            status.update(label="Fertig.", state="complete")
        except Exception as fehler:
            bereinige_verlauf(st.session_state.messages)
            antwort_text = fehlermeldung_fuer(fehler)
            neue_dateien = []
            status.update(label="Fehler.", state="error")

        # Kosten dieser Antwort berechnen und zur Sitzungssumme addieren.
        tokens = verbrauch.get("input", 0) + verbrauch.get("output", 0)
        kosten = geschaetzte_kosten(
            st.session_state.modell, verbrauch.get("input", 0), verbrauch.get("output", 0)
        )
        st.session_state.kosten_total += kosten
        st.session_state.tokens_total += tokens
        info_text = (
            f"ca. US$ {kosten:.3f} · {tokens:,} Tokens" if tokens else ""
        )

        platzhalter.markdown(antwort_text)
        for j, dateiname in enumerate(neue_dateien):
            zeige_download_knopf(dateiname, key=f"livedl_{len(st.session_state.anzeige)}_{j}")
        if info_text:
            st.caption(info_text)

    st.session_state.anzeige.append(
        {"role": "assistant", "text": antwort_text, "dateien": neue_dateien, "info": info_text}
    )
    st.rerun()


def main():
    st.set_page_config(page_title="QV-Pruefungsassistent", page_icon="📝", layout="centered")
    ordner_anlegen()

    # Optionaler Passwortschutz (nur aktiv, wenn ein Passwort gesetzt ist).
    if not pruefe_passwort():
        return

    starte_session_state()

    st.title("📝 QV-Pruefungsassistent")
    st.caption(
        "Assistent fuer die Geleitete Fallarbeit (Kaufleute EFZ, HKB B/C/E). "
        "Laeuft lokal oder online. Keine Telemetrie, keine externe Datenbank."
    )

    baue_seitenleiste()

    # Hinweis, wenn das anthropic-Paket fehlt.
    if not ANTHROPIC_VERFUEGBAR:
        st.error(
            "Das Paket 'anthropic' ist nicht installiert. Bitte schliesse die App und "
            "starte sie ueber start.bat (Windows) bzw. start.sh (Mac/Linux) neu."
        )
        return

    # Hinweis, wenn noch kein Schluessel da ist.
    if not st.session_state.api_key:
        st.warning(
            "Es ist noch kein API-Schluessel hinterlegt. Trage ihn links in der "
            "Seitenleiste ein. Einen Schluessel bekommst du auf **console.anthropic.com** "
            "(dafuer ist API-Guthaben noetig). Danach «Verbindung testen» druecken."
        )

    zeige_verlauf()

    # Datei-Upload ist direkt ins Eingabefeld integriert (Bueroklammer-Symbol).
    # accept_file="multiple" erlaubt eine oder mehrere Dateien pro Nachricht.
    eingabe = st.chat_input(
        "Ausgangslage, Teilaufgabe (z. B. «T1») oder Rueckfrage ...",
        accept_file="multiple",
    )
    if eingabe:
        frage, dateien = entpacke_chat_eingabe(eingabe)
        if not st.session_state.api_key:
            st.error(
                "Bitte zuerst links den API-Schluessel eintragen. Einen Schluessel "
                "bekommst du auf console.anthropic.com (API-Guthaben noetig)."
            )
        elif frage.strip() or dateien:
            pfade = speichere_uploads(dateien)
            behandle_eingabe(frage, pfade)


if __name__ == "__main__":
    main()
